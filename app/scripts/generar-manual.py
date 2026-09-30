#!/usr/bin/env python3
"""
Genera el Manual de Usuario en Word a partir de las capturas de la aplicacion.

El manual es visual a proposito: cada pantalla ocupa una pagina, con la imagen
grande y unos pocos pasos numerados debajo. Va dirigido a personas que no han
usado la plataforma, y para ellas una imagen explica mas que un parrafo.

    npm run manual:generar

Antes hay que tener las capturas, que salen de la aplicacion en marcha con
cuentas y un modulo de ejemplo que el propio guion crea y borra:
    npm run dev
    npm run manual:capturar
"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent.parent
CAPTURAS = RAIZ / "manual" / "capturas"
SALIDA = RAIZ / "MANUAL-DE-USUARIO.docx"

TERRACOTA = RGBColor(0xB8, 0x48, 0x2A)
VERDE = RGBColor(0x17, 0x75, 0x75)
GRIS = RGBColor(0x5A, 0x52, 0x4C)
TINTA = RGBColor(0x2D, 0x1E, 0x16)
FUENTE = "Arial"

ANCHO_IMAGEN = Cm(16.5)

figura = 0


# --------------------------------------------------------------------- Word

def sombrear(celda, color):
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:fill"), color)
    celda._tc.get_or_add_tcPr().append(s)


def campo(parrafo, instruccion):
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    texto = OxmlElement("w:instrText")
    texto.set(qn("xml:space"), "preserve")
    texto.text = instruccion
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    corrida = parrafo.add_run()._r
    for e in (inicio, texto, fin):
        corrida.append(e)


def preparar(doc):
    s = doc.sections[0]
    s.orientation = WD_ORIENT.PORTRAIT
    s.top_margin = s.bottom_margin = Cm(2)
    s.left_margin = s.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    pie = s.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pie.add_run("Manual de Usuario · Memoria El Salado · ")
    campo(pie, "PAGE")
    for c in pie.runs:
        c.font.name = FUENTE
        c.font.size = Pt(8)
        c.font.color.rgb = GRIS


def salto(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def titulo(doc, texto, nivel=2):
    p = doc.add_paragraph(style=f"Heading {nivel}")
    p.paragraph_format.space_before = Pt(0 if nivel == 2 else 12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    c = p.add_run(texto)
    c.bold = True
    c.font.name = FUENTE
    c.font.size = Pt(17 if nivel == 2 else 13)
    c.font.color.rgb = TERRACOTA if nivel == 2 else TINTA
    return p


def parrafo(doc, trozos, tamano=11, color=None, espacio=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(espacio)
    for texto, negrita in trozos:
        c = p.add_run(texto)
        c.bold = negrita
        c.font.name = FUENTE
        c.font.size = Pt(tamano)
        if color:
            c.font.color.rgb = color
    return p


def imagen(doc, archivo, pie):
    """Coloca una captura a lo ancho de la caja de texto, con su pie."""
    global figura
    # Las capturas van numeradas por orden de toma; aqui se buscan por su nombre.
    encontradas = sorted(CAPTURAS.glob(f"[0-9][0-9]-{archivo}.png"))
    if not encontradas:
        raise SystemExit(f"Falta la captura: {archivo}")
    ruta = encontradas[0]

    figura += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ruta), width=ANCHO_IMAGEN)

    pp = doc.add_paragraph()
    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.space_after = Pt(8)
    a = pp.add_run(f"Figura {figura}. ")
    a.bold = True
    b = pp.add_run(pie)
    for c in (a, b):
        c.font.name = FUENTE
        c.font.size = Pt(9)
        c.font.color.rgb = GRIS


def pasos(doc, lista):
    """Pasos numerados: una linea cada uno, con la accion en negrita."""
    for i, item in enumerate(lista, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
        p.paragraph_format.space_after = Pt(3)

        n = p.add_run(f"{i}.  ")
        n.bold = True
        n.font.color.rgb = TERRACOTA
        n.font.name = FUENTE
        n.font.size = Pt(11)

        for texto, negrita in item:
            c = p.add_run(texto)
            c.bold = negrita
            c.font.name = FUENTE
            c.font.size = Pt(11)


def aviso(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    a = p.add_run("⚠  ")
    a.bold = True
    a.font.color.rgb = TERRACOTA
    b = p.add_run(texto)
    for c in (a, b):
        c.font.name = FUENTE
        c.font.size = Pt(10)
        c.font.color.rgb = GRIS


def tabla(doc, filas, anchos=None):
    t = doc.add_table(rows=1, cols=len(filas[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    for celda, texto in zip(t.rows[0].cells, filas[0]):
        celda.text = ""
        sombrear(celda, "F5EDE7")
        c = celda.paragraphs[0].add_run(texto)
        c.bold = True
        c.font.name = FUENTE
        c.font.size = Pt(9.5)

    for fila in filas[1:]:
        celdas = t.add_row().cells
        for celda, texto in zip(celdas, fila):
            celda.text = ""
            c = celda.paragraphs[0].add_run(texto)
            c.font.name = FUENTE
            c.font.size = Pt(9.5)

    for fila in t.rows:
        for celda in fila.cells:
            celda.paragraphs[0].paragraph_format.space_after = Pt(2)
    if anchos:
        for fila in t.rows:
            for celda, ancho in zip(fila.cells, anchos):
                celda.width = ancho

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def pantalla(doc, encabezado, archivo, pie, lista, nota=None, nivel=2):
    """Una pantalla del sistema: titulo, captura y pasos. Ocupa una pagina."""
    titulo(doc, encabezado, nivel)
    imagen(doc, archivo, pie)
    if lista:
        pasos(doc, lista)
    if nota:
        aviso(doc, nota)
    salto(doc)


# ------------------------------------------------------------------ Portada

def portada(doc):
    for _ in range(2):
        doc.add_paragraph()

    for texto, tam, color, neg in [
        ("MANUAL DE USUARIO", 13, GRIS, True),
        ("Memoria El Salado", 36, TERRACOTA, True),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        c = p.add_run(texto)
        c.bold = neg
        c.font.name = FUENTE
        c.font.size = Pt(tam)
        c.font.color.rgb = color

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    c = p.add_run("Guía visual paso a paso")
    c.font.name = FUENTE
    c.font.size = Pt(13)
    c.font.color.rgb = GRIS

    imagen(doc, "portada", "La portada pública de la plataforma.")

    for texto in [
        "Trabajo de grado — Programa de Ingeniería de Software",
        "Universidad Manuela Beltrán",
        "Fabian Santiago Tobón Valero · Daniel Alejandro León Ladino",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        c = p.add_run(texto)
        c.bold = texto.startswith("Universidad")
        c.font.name = FUENTE
        c.font.size = Pt(10.5)
        c.font.color.rgb = GRIS

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    c = p.add_run("Versión 2.0 · Septiembre de 2026")
    c.font.name = FUENTE
    c.font.size = Pt(9.5)
    c.font.color.rgb = GRIS

    salto(doc)


def b(texto):
    """Un trozo en negrita, para los pasos."""
    return (texto, True)


def t(texto):
    """Un trozo normal."""
    return (texto, False)


# --------------------------------------------------------------------- Guia

def construir():
    doc = Document()
    preparar(doc)
    portada(doc)

    # ------------------------------------------------------------ Qué es
    titulo(doc, "Qué es esta plataforma")
    parrafo(doc, [
        t("Memoria El Salado es una plataforma para trabajar en clase la memoria histórica del "
          "conflicto armado colombiano. El contenido no viene escrito dentro del programa: se "
          "construye a partir de las guías del "),
        b("Centro Nacional de Memoria Histórica"),
        t(" en PDF."),
    ])

    titulo(doc, "Quién puede hacer qué", 3)
    tabla(doc, [
        ["Perfil", "Qué hace", "A quién da de alta"],
        ["Administrador", "Importa las guías, arma los módulos y elige la actividad de cada sesión", "A los docentes"],
        ["Docente", "Organiza cursos, publica sesiones, plantea actividades y comenta el trabajo", "A sus estudiantes"],
        ["Estudiante", "Recorre la ruta, hace las actividades y conversa con su curso", "A nadie"],
    ], anchos=[Cm(3.6), Cm(8.4), Cm(4.5)])

    aviso(doc, "Nadie se registra por su cuenta y la plataforma no pide correo. Cada persona recibe "
               "su usuario de quien está un nivel por encima.")

    titulo(doc, "Cómo está organizado el contenido", 3)
    parrafo(doc, [
        t("Un "), b("módulo"), t(" (por ejemplo, El Salado) se divide en "), b("ejes"),
        t(", cada eje en "), b("sesiones"), t(", y cada sesión tiene "), b("una actividad"),
        t(". El docente agrupa a sus estudiantes en "), b("cursos"),
        t(": lo que un estudiante comparte lo ven su curso y su docente."),
    ])
    salto(doc)

    # ---------------------------------------------------------- Entrar
    titulo(doc, "1.  Entrar a la plataforma")
    imagen(doc, "login", "Pantalla de entrada, con el rol elegido.")
    pasos(doc, [
        [t("En la portada, pulsa "), b("Entrar"), t(" o "), b("Comenzar viaje"), t(".")],
        [t("Pulsa tu "), b("rol"), t(": Estudiante, Docente o Administrador.")],
        [t("Escribe tu "), b("usuario"), t(" y tu "), b("contraseña"), t(". El ojo deja verla mientras la escribes.")],
        [t("Pulsa "), b("Entrar a la Plataforma"), t(".")],
    ])
    aviso(doc, "Si olvidaste tu contraseña, pulsa «¿Olvidaste tu contraseña?»: te dice a quién pedirla. "
               "Al estudiante se la restablece su docente; al docente, la administración.")
    salto(doc)

    pantalla(
        doc,
        "2.  La primera vez: elige tu contraseña",
        "primer-ingreso",
        "Hasta que elijas tu contraseña, la plataforma no te deja pasar a ninguna otra pantalla.",
        [
            [t("Escribe la "), b("contraseña provisional"), t(" que te entregaron.")],
            [t("Escribe la "), b("nueva"), t(" y "), b("repítela"), t(" debajo.")],
            [t("Mira la lista: te dice qué le falta a tu contraseña.")],
            [t("Cuando no quede ningún aviso, pulsa "), b("Guardar y continuar"), t(".")],
        ],
        "Tu contraseña necesita 10 caracteres o más, con mayúscula, minúscula, número y símbolo. "
        "No puede llevar tu nombre ni tu usuario.",
    )

    pantalla(
        doc,
        "3.  Estudiante: elige tu explorador",
        "explorador",
        "En la primera entrada, el estudiante elige quién lo acompaña en el viaje.",
        [
            [t("Elige uno de los cuatro exploradores: "), b("Curioso, Investigador, Guardián o Semillero"), t(".")],
            [t("Pulsa "), b("Continuar al mapa"), t(".")],
            [t("Tu explorador aparece en la cabecera. Lo puedes cambiar desde "), b("Cambiar mi explorador"), t(", al pie del menú.")],
        ],
    )

    # ---------------------------------------------------------- Estudiante
    titulo(doc, "Guía del estudiante")
    parrafo(doc, [
        t("Tu menú tiene el "), b("Mapa del Viaje"), t(", desde donde entras a todo, los "),
        b("Debates"), t(" de tu curso y "), b("Mi Diario"), t("."),
    ])
    imagen(doc, "estudiante-mapa", "El Mapa del Viaje: los ejes, tu avance y tus competencias.")
    pasos(doc, [
        [t("Arriba ves tu avance y las cuatro "), b("competencias"), t(" de la guía: empatía, escucha activa, pensamiento crítico y asertividad.")],
        [t("Pulsa una tarjeta de eje para ver sus sesiones.")],
        [t("Si tu docente abre los ejes en orden, los siguientes aparecen con candado hasta que terminas el anterior.")],
    ])
    salto(doc)

    pantalla(
        doc,
        "Abrir una sesión",
        "estudiante-mapa-eje",
        "Al desplegar un eje aparecen sus sesiones.",
        [
            [t("Las sesiones "), b("publicadas"), t(" tienen el botón "), b("Abrir"), t(".")],
            [t("Las que están en gris aún no las habilitó tu docente.")],
        ],
        "Cada sesión lleva su propia actividad y su propio texto, tomado de la guía del CNMH.",
        nivel=3,
    )

    pantalla(
        doc,
        "Cartografía social",
        "estudiante-cartografia",
        "Marca en un mapa real los lugares que tienen significado para ti.",
        [
            [t("Haz clic en el "), b("mapa"), t(" sobre el lugar que quieres marcar.")],
            [t("Escribe el "), b("nombre del lugar"), t(" y tu "), b("relato"), t(". Si prefieres, "), b("grábalo con tu voz"), t(".")],
            [t("Añade una "), b("foto"), t(" del lugar si la tienes, y elige las sensaciones que te evoca.")],
            [t("Pulsa "), b("Sembrar en el mapa"), t(". Tu marca la ven tu curso y tu docente; si tu docente armó grupos, puedes marcar en nombre del tuyo.")],
        ],
        nivel=3,
    )

    pantalla(
        doc,
        "La despensa del Caribe (copla)",
        "estudiante-despensa",
        "Ubica en el mapa de los Montes de María lo que produce la región.",
        [
            [t("Toca un "), b("producto"), t(" y luego el "), b("lugar"), t(" donde se da. También puedes arrastrarlo.")],
            [t("Explica por qué la región se llamaba «la despensa del Caribe» y, si quieres, escribe tus versos.")],
            [t("Pulsa "), b("Guardar"), t(": después verás cómo lo ubicó tu curso.")],
        ],
        nivel=3,
    )

    pantalla(
        doc,
        "Archivo y anotaciones",
        "estudiante-anotaciones",
        "El texto de la sesión y, al lado, la conversación de tu curso.",
        [
            [t("Haz "), b("clic sobre un párrafo"), t(" para citarlo, escribe tu reflexión y pulsa "), b("Publicar anotación"), t(".")],
            [t("Pulsa "), b("Responder"), t(" para contestar la anotación de un compañero, y el pulgar si "), b("te sirve"), t(".")],
            [t("En "), b("Destacados"), t(" están las anotaciones que tu docente marcó para la discusión.")],
        ],
        nivel=3,
    )

    pantalla(
        doc,
        "Simulador de repartición de tierras",
        "estudiante-simulador",
        "Repartes los 20 lotes de la finca «La Esperanza» y $100.000 en subsidios entre 20 familias.",
        [
            [b("Arrastra"), t(" una familia a un lote, o pulsa "), b("Elegir"), t(" en su tarjeta y toca el lote.")],
            [t("También puedes tocar un lote y escoger en la lista a qué familia dárselo, o dejarlo libre.")],
            [t("Elige el "), b("subsidio"), t(" de cada familia en su tarjeta. Cada lote dice si la tierra alcanza y si el subsidio es suficiente.")],
            [t("Vigila la "), b("equidad"), t(" y el "), b("riesgo"), t(", y pulsa "), b("Validar distribución"), t(" para guardar.")],
        ],
        "No hay una única respuesta correcta. El ejercicio sirve para ver las tensiones entre "
        "atender a todos y atender bien a quien más lo necesita. Una familia puede tener varios lotes, hasta 40 ha.",
        nivel=3,
    )

    pantalla(
        doc,
        "Dilema",
        "estudiante-dilema",
        "Tu docente plantea el dilema; tú eliges y lo justificas.",
        [
            [t("Lee la pregunta y "), b("elige una opción"), t(".")],
            [t("Escribe por qué y pulsa "), b("Guardar"), t(".")],
            [t("Después verás "), b("cómo eligió tu curso"), t(", y el debate que tu docente haya abierto sobre la sesión.")],
        ],
        "Me ubico, A investigar, Testimonios y Compromiso funcionan igual: una pregunta, tu respuesta y, "
        "cuando corresponde, lo que compartió tu curso.",
        nivel=3,
    )

    pantalla(
        doc,
        "Alerta de cuidado",
        "estudiante-cuidado",
        "Antes de los contenidos más duros, la plataforma avisa y ofrece acompañamiento.",
        [
            [t("Lee el aviso antes de empezar.")],
            [t("Pulsa "), b("Entiendo, continuar"), t(" cuando te sientas listo, o "), b("Dejarlo para después"), t(".")],
            [t("Si prefieres hablar antes, pulsa "), b("Hablar con mi docente"), t(": el pedido le llega a él, no a tu curso.")],
        ],
        "El botón «Alerta de cuidado» del menú está siempre disponible.",
        nivel=3,
    )

    pantalla(
        doc,
        "Jardín de la Memoria",
        "estudiante-jardin",
        "Cada flor recuerda a una persona o un lugar, con una carta, un poema, un dibujo o una foto.",
        [
            [t("Pulsa "), b("Sembrar una Memoria"), t(" y escribe a quién o qué recuerdas, y su legado.")],
            [t("Elige qué dejar: "), b("una carta, un poema, un dibujo o la foto de una obra"), t(".")],
            [t("Pulsa "), b("Sembrar en el jardín"), t(". La flor aparece en el jardín de tu curso con tu nombre.")],
        ],
        nivel=3,
    )

    pantalla(
        doc,
        "Compromiso",
        "estudiante-compromiso",
        "Al cerrar la ruta, cada estudiante escribe a qué se compromete.",
        [
            [t("Escribe tu compromiso y pulsa "), b("Guardar"), t(".")],
            [t("Si marcas "), b("Compartirlo con mi curso"), t(", aparece con tu nombre en los "), b("compromisos de tu curso"), t(".")],
        ],
        nivel=3,
    )

    pantalla(
        doc,
        "Avisos",
        "estudiante-avisos",
        "La campana de la cabecera avisa de los comentarios de tu docente y de las sesiones nuevas.",
        [
            [t("El número sobre la campana dice cuántos avisos nuevos tienes.")],
            [t("Pulsa la campana para verlos, y el aviso para ir a donde está.")],
        ],
        nivel=3,
    )

    pantalla(
        doc,
        "Mi Diario",
        "estudiante-diario",
        "Espacio personal. Tú decides qué se comparte.",
        [
            [t("Elige "), b("cómo te sientes"), t(".")],
            [t("Escribe lo que quieras recordar.")],
            [t("Marca "), b("compartir"), t(" solo si quieres que tu docente lo lea. Sus comentarios aparecen debajo de la entrada.")],
        ],
        "Tus entradas son privadas por defecto. Nadie puede leer una entrada que no hayas "
        "marcado como compartida.",
        nivel=3,
    )

    # ------------------------------------------------------------- Docente
    titulo(doc, "Guía del docente")
    imagen(doc, "docente-panel", "Panel del docente: avance, competencias, sesiones y pedidos de acompañamiento.")
    pasos(doc, [
        [t("Arriba eliges el "), b("curso"), t(" que quieres ver.")],
        [t("En "), b("Gestión de Sesiones"), t(", el interruptor publica o retira cada sesión, y el enlace de cada una lleva a su página.")],
        [t("A la derecha llegan los "), b("pedidos de acompañamiento"), t(" y los trabajos por revisar.")],
    ])
    salto(doc)

    pantalla(
        doc,
        "Cursos y grupos",
        "docente-cursos",
        "Tus cursos, con su año lectivo.",
        [
            [t("Escribe el nombre y el año, y pulsa "), b("Crear curso"), t(".")],
            [t("Abre un curso para asignarle estudiantes y armar grupos de trabajo.")],
        ],
        nivel=3,
    )

    pantalla(
        doc,
        "El recorrido de un curso",
        "docente-curso",
        "Dentro del curso decides si los ejes se recorren libres o en orden.",
        [
            [t("Suma estudiantes con "), b("Sumar"), t(" y arma grupos de trabajo con "), b("Crear grupo"), t(".")],
            [t("En "), b("Recorrido del mapa"), t(", elige "), b("Libre"), t(" o "), b("En orden, eje por eje"), t(", y pulsa "), b("Guardar recorrido"), t(".")],
            [t("En orden, puedes abrir un eje para todo el curso antes de tiempo con "), b("Abrir ya"), t(".")],
        ],
        nivel=3,
    )

    pantalla(
        doc,
        "Dar de alta a tus estudiantes",
        "docente-estudiantes",
        "El formulario ya propone una contraseña provisional.",
        [
            [t("Escribe el "), b("nombre"), t(" y el "), b("usuario"), t(", y elige su curso.")],
            [t("Pulsa "), b("Crear estudiante"), t(" y entrégale las credenciales.")],
            [t("Si el acudiente lo autoriza, registra su "), b("subpoblación"), t(" con el consentimiento.")],
            [t("Si alguien olvida su contraseña, usa "), b("Restablecer contraseña"), t(".")],
        ],
        "La contraseña provisional no se vuelve a mostrar. Anótala antes de salir de la pantalla.",
        nivel=3,
    )

    pantalla(
        doc,
        "Revisión",
        "docente-revision",
        "Las respuestas, las simulaciones validadas y los diarios compartidos de tus estudiantes.",
        [
            [t("Filtra por "), b("curso"), t(" o por "), b("pendientes"), t(".")],
            [t("Escribe tu comentario y pulsa "), b("Comentar"), t(". Al estudiante le llega en la campana.")],
        ],
        nivel=3,
    )

    pantalla(
        doc,
        "Las anotaciones de una sesión",
        "docente-anotaciones",
        "La conversación de tus estudiantes sobre el texto de la sesión.",
        [
            [t("Filtra por "), b("curso"), t(" si tienes varios.")],
            [t("Pulsa "), b("Responder"), t(" para entrar en un hilo, o anota tú también.")],
            [t("Pulsa "), b("Destacar"), t(" en las que ayudan a todo el curso: aparecen en su pestaña Destacados.")],
        ],
        nivel=3,
    )

    pantalla(
        doc,
        "Plantear un dilema",
        "docente-dilema",
        "El dilema lo escribe cada docente, adaptado a su contexto.",
        [
            [t("Escribe la pregunta y entre dos y cuatro "), b("opciones"), t(".")],
            [t("Pulsa "), b("Plantear el dilema"), t(". Debajo verás cómo eligieron tus estudiantes.")],
        ],
        "En «A investigar» se plantea igual la pregunta de investigación.",
        nivel=3,
    )

    pantalla(
        doc,
        "Contenidos y recursos",
        "docente-contenidos",
        "Las guías del CNMH, incluida la de maestros que el estudiante no puede abrir.",
        [
            [t("Arriba están las "), b("guías del módulo activo"), t(".")],
            [t("Debajo subes "), b("recursos"), t(": audio, video, imagen o PDF, o un video de YouTube o Vimeo.")],
            [t("Los audios y videos piden "), b("transcripción"), t(" o subtítulos antes de publicarse.")],
        ],
        nivel=3,
    )

    pantalla(
        doc,
        "Reportes por subpoblaciones",
        "docente-subpoblaciones",
        "Avance comparado según el vínculo de cada estudiante con el conflicto, solo con consentimiento.",
        [
            [t("Revisa el "), b("avance"), t(" de cada subpoblación.")],
            [t("Pulsa "), b("Generar reporte diferenciado"), t(" para dejarlo por escrito.")],
        ],
        "Los grupos de menos de cinco estudiantes no se muestran, para que nadie pueda ser identificado. "
        "No es una calificación.",
        nivel=3,
    )

    pantalla(
        doc,
        "Debates y foros",
        "docente-foros",
        "Debate estructurado en dos roles contrapuestos, con argumentos y fuentes.",
        [
            [t("Escribe la "), b("pregunta orientadora"), t(" y los dos "), b("roles"), t(".")],
            [t("Si quieres, asócialo a una sesión: aparecerá dentro de ella.")],
            [t("Pulsa "), b("Abrir foro"), t(". Puedes cerrarlo cuando termine el debate.")],
        ],
        nivel=3,
    )

    # ------------------------------------------------------- Administrador
    titulo(doc, "Guía del administrador")
    imagen(doc, "admin-docentes", "Alta de docentes y entrega de credenciales.")
    pasos(doc, [
        [t("Escribe el "), b("nombre"), t(" y el "), b("usuario"), t(" del docente.")],
        [t("Acepta la contraseña propuesta o pulsa "), b("Generar otra"), t(".")],
        [t("Pulsa "), b("Crear docente"), t(" y copia las credenciales.")],
        [t("Si alguien la pierde, usa "), b("Restablecer contraseña"), t(".")],
    ])
    salto(doc)

    pantalla(
        doc,
        "Importar una guía en PDF",
        "admin-importar",
        "Los documentos cargados, con su módulo y su visibilidad.",
        [
            [t("Sube el PDF y pulsa "), b("Analizar documento"), t(".")],
            [t("Revisa la estructura: títulos, actividad y "), b("texto de cada sesión"), t(".")],
            [t("Completa a mano el texto que no se haya podido extraer.")],
            [t("Da nombre al módulo y créalo.")],
        ],
        nivel=3,
    )

    titulo(doc, "Las dos guías del CNMH", 3)
    parrafo(doc, [
        t("Cada caso se publica en dos documentos, y no cumplen la misma función:"),
    ])
    tabla(doc, [
        ["Documento", "Para qué sirve", "Quién puede abrirlo"],
        ["Guía del estudiante\n(el-salado.pdf)", "Construir el módulo: trae la ruta y las actividades",
         "Todos"],
        ["Guía para maestros\n(el-salado_guia-para-maestros.pdf)",
         "Preparar la clase: trae las orientaciones y lo que se espera que descubran",
         "Solo el docente"],
    ], anchos=[Cm(5.2), Cm(7.5), Cm(3.8)])

    pasos(doc, [
        [t("Importa la "), b("guía del estudiante"), t(" y crea el módulo con ella.")],
        [t("Sube la "), b("guía para maestros"), t(".")],
        [t("Pulsa "), b("Reservar al docente"), t(" sobre ella.")],
        [t("En su selector, elige el módulo para asociarla.")],
    ])
    aviso(doc, "Si el estudiante puede abrir la guía para maestros, la actividad pierde sentido: "
               "le entrega de antemano las respuestas a las que debía llegar solo.")
    salto(doc)

    pantalla(
        doc,
        "Módulos geográficos",
        "admin-casos",
        "Los casos de estudio creados, ubicados en el mapa de Colombia.",
        [
            [t("Consulta el "), b("estado"), t(" y el número de sesiones de cada módulo.")],
            [t("Crea un caso nuevo heredando la estructura del activo.")],
            [t("Abre un módulo para cambiar la actividad de sus sesiones.")],
        ],
        "Eliminar un módulo borra también el trabajo que los estudiantes hicieron en sus "
        "sesiones. El PDF de origen no se toca.",
        nivel=3,
    )

    pantalla(
        doc,
        "La actividad de cada sesión",
        "admin-actividades",
        "Cada sesión del módulo con su actividad y la que sugiere el importador.",
        [
            [t("Cambia la actividad de una sesión en su selector, o pulsa "), b("Usar la sugerida"), t(".")],
            [t("Pulsa "), b("Aplicar todas las sugeridas"), t(" para usar las del importador en todo el módulo.")],
        ],
        "Cambiar la actividad de una sesión que ya se trabajó no borra lo que los estudiantes hicieron, "
        "pero la pantalla que verán será la nueva.",
        nivel=3,
    )

    pantalla(
        doc,
        "Auditoría de accesibilidad",
        "admin-accesibilidad",
        "Transcripciones, subtítulos y contraste de los recursos, calculados a partir de los datos.",
        [
            [t("Revisa el "), b("cumplimiento global"), t(" y el detalle por recurso.")],
            [t("Los recursos sin transcripción aparecen marcados para corregir.")],
        ],
        nivel=3,
    )

    # ------------------------------------------------------- Cosas comunes
    titulo(doc, "Cambiar tu contraseña cuando quieras")
    imagen(doc, "cambiar-contrasena", "Se llega desde «Cambiar mi contraseña», al pie del menú.")
    pasos(doc, [
        [t("Escribe tu "), b("contraseña actual"), t(".")],
        [t("Escribe la "), b("nueva"), t(" dos veces.")],
        [t("Pulsa "), b("Guardar y continuar"), t(".")],
    ])
    salto(doc)

    # ------------------------------------------------------ Si algo falla
    titulo(doc, "Si algo no funciona")
    tabla(doc, [
        ["Lo que ves", "Por qué pasa", "Qué hacer"],
        ["«Credenciales incorrectas o el rol no corresponde»",
         "El rol elegido no es el tuyo, o hay un error al escribir",
         "Revisa primero el rol; es lo más común"],
        ["Vuelve siempre a «Elige tu contraseña»",
         "Todavía usas la contraseña provisional",
         "Complétala: es obligatorio"],
        ["No veo ninguna sesión",
         "El docente no las ha publicado",
         "Pídele que las habilite en su panel"],
        ["Un eje aparece con candado",
         "Tu curso recorre los ejes en orden",
         "Termina el anterior o pregúntale a tu docente"],
        ["«El documento no traía texto para esta sesión»",
         "El PDF no desarrollaba esa parte",
         "El administrador puede pegarlo al importar"],
        ["El mapa se ve como una cuadrícula",
         "No hay conexión a internet",
         "Nada: la actividad funciona igual"],
        ["Olvidé mi contraseña",
         "No se puede recuperar, está cifrada",
         "Pide una provisional nueva a tu docente, o a la administración si eres docente"],
    ], anchos=[Cm(5.2), Cm(5.5), Cm(5.8)])

    titulo(doc, "Antes de cerrar", 3)
    parrafo(doc, [
        t("Pulsa "), b("Salir"),
        t(" en la esquina superior derecha, sobre todo si el computador es compartido."),
    ])

    doc.save(SALIDA)
    print(f"Manual generado: {SALIDA.name} · {figura} figuras")


if __name__ == "__main__":
    construir()
